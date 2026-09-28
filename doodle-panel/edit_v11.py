import copy
from pptx import Presentation
from pptx.util import Emu, Inches

SRC = 'in.pptx'
OUT = 'Doodle_Labs_Europe_Panel_v11.pptx'
p = Presentation(SRC)


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_lines(sh, lines):
    """Replace text paragraph by paragraph, keeping the first run's formatting."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):  # clone last paragraph
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for para, line in zip(paras, lines):
        runs = para.runs
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for para in paras[len(lines):]:
        para._p.getparent().remove(para._p)


def delete(slide, sid):
    sh = shape(slide, sid)
    sh._element.getparent().remove(sh._element)


def move_y(slide, sid, y_in):
    shape(slide, sid).top = Emu(int(y_in * 914400))


def clone_below(slide, sid, y_in, lines, bold=None):
    src = shape(slide, sid)
    el = copy.deepcopy(src._element)
    src._element.addnext(el)
    # new shape is the element right after src; give it a unique id
    new_id = max(s.shape_id for s in slide.shapes) + 1
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('id', str(new_id))
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('name', 'Added text %d' % new_id)
    new = shape(slide, new_id)
    new.top = Emu(int(y_in * 914400))
    set_lines(new, lines)
    if bold is not None:
        for para in new.text_frame.paragraphs:
            for r in para.runs:
                r.font.bold = bold
    return new


def note_replace(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            para.runs[0].text = full.replace(old, new)
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
            return
    raise ValueError('note text not found: ' + old[:50])



from pptx.dml.color import RGBColor
S = {i: p.slides[i - 1] for i in range(1, len(p.slides) + 1)}
RED = shape(S[4], 17).text_frame.paragraphs[0].runs[0].font.color.rgb
DARK = RGBColor(0x20, 0x21, 0x24)

# ---- 03 · Where I would start -> own Southern Europe ----
s = S[4]
set_lines(shape(s, 19), ['Market  /  Where I would start'])
set_lines(shape(s, 3), ['I would own Southern Europe,', 'starting where a decision is reachable'])
set_lines(shape(s, 15), ['Build through the primes;', 'programmes from 2028'])
set_lines(shape(s, 17), ['One owner for the South. A live decision in your book goes to the top.'])
note_replace(s, 'I rank by where a radio decision can actually be reached, not by the biggest order or by my languages.',
             'I would take ownership of Southern Europe — France, Italy, Spain and Portugal — and sequence it by where a radio decision can actually be reached, not by the biggest order or by my languages.')
note_replace(s, 'Spain and Portugal I would qualify rather than lead with, even though they are my home ground:',
             'Spain and Portugal I would build through the primes, even though they are my home ground:')

# ---- 04 · Where Doodle fits -> the play travels ----
s = S[5]
set_lines(shape(s, 25), ['Market  /  Where Doodle fits'])
set_lines(shape(s, 3), ['An OEM designs the link in once:', 'Sky Mantis 2 shows how'])
clone_below(s, 22, 4.65, ['Travels'])
clone_below(s, 23, 4.65, ['Works in any country: decisions sit with the platform, not the border.'])
note_replace(s, 'I do not promise range or approvals; Sales Engineering and the customer agree the test in their configuration.',
             'I do not promise range or approvals; Sales Engineering and the customer agree the test in their configuration. And the play travels: Evolve is a UK OEM. The question is about the platform, not the border — which is why I would also take named accounts outside my core.')

# ---- 05 · New territory -> any country or named account ----
s = S[6]
set_lines(shape(s, 17), ['New territory'])
set_lines(shape(s, 3), ['Any new country or named account:'])
_r = shape(s, 3).text_frame.paragraphs[0].runs[0]._r
_br = _r.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}br', {})
_rpr = _r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
if _rpr is not None: _br.append(copy.deepcopy(_rpr))
_r2 = copy.deepcopy(_r)
_r2.find('{http://schemas.openxmlformats.org/drawingml/2006/main}t').text = 'the same opening move'
_r.addnext(_br); _br.addnext(_r2)
set_lines(shape(s, 4), ['“For your next platform, who leads the datalink requirement, and when do you freeze the radio architecture?”'])
for para in shape(s, 4).text_frame.paragraphs:
    for r in para.runs:
        r.font.color.rgb = DARK; r.font.underline = False; r.font.italic = False
set_lines(shape(s, 18), ['How I opened Portugal’s grid operator from zero: a partner who held the relationship.'])
ask = clone_below(s, 15, 5.10, ['Give me a new country or a named account: I run this play and report at day 30.'], bold=True)
for para in ask.text_frame.paragraphs:
    for r in para.runs:
        r.font.color.rgb = RED
note_replace(s, 'Say I am given a country with no Doodle business.',
             'Say I am given a country with no Doodle business, or a single named account outside the South — the play is the same.')
note_replace(s, 'I built the numbers, secured the permit and we won the client’s tender.',
             'I built the numbers, secured the permit and we won the client’s tender. So my ask is simple: give me a new country or a named account, and at day thirty you get the owner, the decision date and my recommendation.')

# ---- 11 · Scope -> consistent with the story ----
s = S[12]
set_lines(shape(s, 3), ['Proposed scope: own Southern Europe,', 'open to named accounts and new countries'])
set_lines(shape(s, 11), ['Named accounts and new countries: yes, where I add access.', 'Germany as Andy decides, one owner per customer.'])
note_replace(s, 'On scope, my proposal is France, Italy and Iberia as a core,',
             'On scope, my proposal is to own Southern Europe — France, Italy and Iberia — as a core,')
note_replace(s, 'Elsewhere I would take named accounts only where I add access or continuity a customer lacks today,',
             'Beyond that, I want to be considered for named accounts and new countries wherever I add access or continuity a customer lacks today,')
note_replace(s, '“Slide 3 says don’t lead with Italy and Spain, yet they are in your core.”',
             '“Slide 3 puts Spain and Portugal third, yet they are in your core.”')
note_replace(s, 'The core is coverage and customer continuity; slide 3 is the order in which I start new outreach.',
             'The core is what I own; slide 3 is only the order in which I start new outreach inside it.')

note_replace(S[4], 'And if your book has a live customer decision, that goes to the top of this list.',
             'And if your book has a live customer decision, that goes to the top of this list. In Spain and Italy the primes are customers, not competitors — the same stance DroneShield takes publicly: we sell inside their platforms.')
note_replace(S[12], 'The customer should see one plan and one Doodle contact.',
             'The customer should see one plan and one Doodle contact. Coverage should follow the type of deal, not only the map — McKinsey’s point about letting the transaction decide the go-to-market model — which is exactly why a named-account overlay makes sense next to a regional core.')
note_replace(S[6], 'So my ask is simple: give me a new country or a named account, and at day thirty you get the owner, the decision date and my recommendation.',
             'So my ask is simple: give me a new country or a named account, and at day thirty you get the owner, the decision date and my recommendation. '
             'IF ASKED “How would you decide which new country to open next?” — Three tests, all needed: external readiness (a funded programme with a radio decision still open), a reference that travels (a customer or case the OEM will recognise), and internal readiness (Sales Engineering capacity and a direct or partner route). When all three are true, I open it; when one is missing, I record the trigger and wait.')
p.save(OUT)
print('saved', OUT)
