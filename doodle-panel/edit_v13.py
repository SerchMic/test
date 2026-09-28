import copy
from pptx import Presentation
from pptx.util import Emu, Inches

SRC = 'Doodle_Labs_Europe_Panel_v12.pptx'
OUT = 'Doodle_Labs_Europe_Panel_v13.pptx'
p = Presentation(SRC)


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_lines(sh, lines):
    """Replace text paragraph by paragraph, keeping the first run's formatting."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):  # clone last paragraph
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for para, line in zip(paras, lines):
        runs = para.runs
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for para in paras[len(lines):]:
        para._p.getparent().remove(para._p)


def delete(slide, sid):
    sh = shape(slide, sid)
    sh._element.getparent().remove(sh._element)


def move_y(slide, sid, y_in):
    shape(slide, sid).top = Emu(int(y_in * 914400))


def clone_below(slide, sid, y_in, lines, bold=None):
    src = shape(slide, sid)
    el = copy.deepcopy(src._element)
    src._element.addnext(el)
    # new shape is the element right after src; give it a unique id
    new_id = max(s.shape_id for s in slide.shapes) + 1
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('id', str(new_id))
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('name', 'Added text %d' % new_id)
    new = shape(slide, new_id)
    new.top = Emu(int(y_in * 914400))
    set_lines(new, lines)
    if bold is not None:
        for para in new.text_frame.paragraphs:
            for r in para.runs:
                r.font.bold = bold
    return new


def note_replace(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            para.runs[0].text = full.replace(old, new)
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
            return
    raise ValueError('note text not found: ' + old[:50])





S = {i: p.slides[i - 1] for i in range(1, len(p.slides) + 1)}

def transition(slide, new):
    tf = slide.notes_slide.notes_text_frame
    paras = list(tf.paragraphs)
    for k, para in enumerate(paras):
        if para.text.strip() == 'TRANSITION':
            nxt = paras[k + 1]
            nxt.runs[0].text = new
            for r in nxt.runs[1:]:
                r._r.getparent().remove(r._r)
            return
    raise ValueError('no TRANSITION')

# --- text changes (old positions) ---
transition(S[2], 'Three funded signals show it.')
transition(S[3], 'First, what an OEM actually buys from Doodle.')
transition(S[5], 'Given that, here is the ground I would want to own.')
transition(S[12], 'Inside that scope, here is where I would start.')
transition(S[4], 'And where Doodle has no business yet, this is the first conversation.')
transition(S[11], 'Which brings me to the real question behind this panel: what you could hold me to.')
note_replace(S[12], '“Slide 3 puts Spain and Portugal third, yet they are in your core.”',
             '“The next slide puts Spain and Portugal third, yet they are in your core.”')
note_replace(S[12], 'slide 3 is only the order in which I start new outreach inside it.',
             'the next slide is only the order in which I start new outreach inside it.')
set_lines(shape(S[4], 3), ['Inside the South, I start', 'where a decision is reachable'])
set_lines(shape(S[4], 19), ['Scope  /  Where I would start'])
for sh in S[11].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == 'Inherit - Input':
        set_lines(sh, ['Inherit'])

# remove stray trailing line break in the new-territory title
t = shape(S[6], 3).text_frame.paragraphs[0]._p
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
kids = [k for k in t if k.tag in (A + 'r', A + 'br')]
while kids and kids[-1].tag == A + 'br':
    t.remove(kids[-1]); kids.pop()

# --- reorder: topic 1 (market, fit) -> scope -> topic 2 -> topic 3 -> topic 4 -> close ---
new_order = [1, 2, 3, 5, 12, 4, 6, 7, 8, 9, 10, 11, 13]
lst = p.slides._sldIdLst
ids = list(lst)
for el in ids:
    lst.remove(el)
for i in new_order:
    lst.append(ids[i - 1])

# --- page labels = position ---
for pos, slide in enumerate(p.slides, 1):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.left > 9.0 * 914400 and sh.top < 0.5 * 914400:
            txt = sh.text_frame.text.strip()
            if txt.isdigit() and len(txt) == 2:
                set_lines(sh, ['%02d' % pos])

p.save(OUT)
print('saved', OUT)
