import copy
from pptx import Presentation
from pptx.util import Emu, Inches

SRC = 'Doodle_Labs_Europe_Panel_v11.pptx'
OUT = 'Doodle_Labs_Europe_Panel_v12.pptx'
p = Presentation(SRC)


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_lines(sh, lines):
    """Replace text paragraph by paragraph, keeping the first run's formatting."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):  # clone last paragraph
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for para, line in zip(paras, lines):
        runs = para.runs
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for para in paras[len(lines):]:
        para._p.getparent().remove(para._p)


def delete(slide, sid):
    sh = shape(slide, sid)
    sh._element.getparent().remove(sh._element)


def move_y(slide, sid, y_in):
    shape(slide, sid).top = Emu(int(y_in * 914400))


def clone_below(slide, sid, y_in, lines, bold=None):
    src = shape(slide, sid)
    el = copy.deepcopy(src._element)
    src._element.addnext(el)
    # new shape is the element right after src; give it a unique id
    new_id = max(s.shape_id for s in slide.shapes) + 1
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('id', str(new_id))
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('name', 'Added text %d' % new_id)
    new = shape(slide, new_id)
    new.top = Emu(int(y_in * 914400))
    set_lines(new, lines)
    if bold is not None:
        for para in new.text_frame.paragraphs:
            for r in para.runs:
                r.font.bold = bold
    return new


def note_replace(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            para.runs[0].text = full.replace(old, new)
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
            return
    raise ValueError('note text not found: ' + old[:50])




S = {i: p.slides[i - 1] for i in range(1, len(p.slides) + 1)}
def by_text(slide, start):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(start):
            return sh
    raise KeyError(start)

# 1 · Title: hunter first, mirrors "What Success Looks Like"; fix the missing "No"
s = S[1]
set_lines(shape(s, 2), ['Win new drone OEMs. Expand each one across its next platforms.'])
set_lines(shape(s, 5), ['Built on public sources only. No Doodle Labs CRM, pricing or quota used.'])
note_replace(s, 'My proposal is simple: protect the customers Doodle already has, and win the next platform decision at OEMs that are moving into production.',
             'My proposal is simple: win new drone OEMs, expand each one across its next platforms and programmes, and protect the customers Doodle already has.')

# 05 · the ask, in the JD's words
set_lines(by_text(S[6], 'Give me a new country'), ['Anywhere in Europe, give me a priority OEM account: I run this play and report at day 30.'])

# Day 1–30: the JD's own success metric
s = S[9]
set_lines(shape(s, 29), ['First customer meetings secured'])
set_lines(shape(s, 33), ['~20 priorities with owners, and first customer meetings secured.'])
note_replace(s, 'The first month is not a research phase.',
             'The first month is not a research phase — your own job description asks for initial customer meetings secured within the first 30 days, and that is inside this plan, not after it.')

# Day 31–60: consistent labels; brand and presence in the notes
s = S[10]
set_lines(shape(s, 13), ['Inherit'])
set_lines(shape(s, 15), ['Do'])
note_replace(s, 'In month two I would convert qualified interest into commitments.',
             'In month two I would convert qualified interest into commitments, and start building Doodle’s presence in the region: the March 2027 shows — Enforce Tac, DSEI Germany, XPONENTIAL Europe — planned with customer meetings set in advance, not booth time.')

# Scope: account-first across Europe, the South as the base
s = S[12]
set_lines(shape(s, 3), ['Own Southern Europe;', 'hunt priority OEM accounts across Europe'])
set_lines(shape(s, 10), ['ACROSS EUROPE'])
set_lines(shape(s, 11), ['Priority drone OEM accounts anywhere, where I add access.', 'Germany as Andy decides, one owner per customer.'])
note_replace(s, 'On scope, my proposal is to own Southern Europe',
             'Your posting describes this role as driving revenue across key drone OEM accounts throughout Europe, and Ryan described it as owning the EU region. So my proposal is account-first: I hunt priority OEM accounts across Europe, and I own Southern Europe')

p.save(OUT)
print('saved', OUT)
