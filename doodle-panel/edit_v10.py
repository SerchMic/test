import copy
from pptx import Presentation
from pptx.util import Emu, Inches

SRC = 'v9.pptx'
OUT = 'Doodle_Labs_Europe_Panel_v10.pptx'
p = Presentation(SRC)


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_lines(sh, lines):
    """Replace text paragraph by paragraph, keeping the first run's formatting."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):  # clone last paragraph
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for para, line in zip(paras, lines):
        runs = para.runs
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for para in paras[len(lines):]:
        para._p.getparent().remove(para._p)


def delete(slide, sid):
    sh = shape(slide, sid)
    sh._element.getparent().remove(sh._element)


def move_y(slide, sid, y_in):
    shape(slide, sid).top = Emu(int(y_in * 914400))


def clone_below(slide, sid, y_in, lines, bold=None):
    src = shape(slide, sid)
    el = copy.deepcopy(src._element)
    src._element.addnext(el)
    # new shape is the element right after src; give it a unique id
    new_id = max(s.shape_id for s in slide.shapes) + 1
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('id', str(new_id))
    el.xpath('./p:nvSpPr/p:cNvPr')[0].set('name', 'Added text %d' % new_id)
    new = shape(slide, new_id)
    new.top = Emu(int(y_in * 914400))
    set_lines(new, lines)
    if bold is not None:
        for para in new.text_frame.paragraphs:
            for r in para.runs:
                r.font.bold = bold
    return new


def note_replace(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            para.runs[0].text = full.replace(old, new)
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
            return
    raise ValueError('note text not found: ' + old[:50])


S = {i: p.slides[i - 1] for i in range(1, 13)}

# ---------- Slide 2: notes — Auterion bridge ----------
note_replace(S[2], 'So the useful question is which upcoming platform still has a communications decision to make.',
             'So the useful question is which upcoming platform still has a communications decision to make. '
             'One of these three, Auterion, already names Doodle on its own website — a public link, not a claim that it is a customer.')

# ---------- Slide 3: Italy second, Portugal folded into third ----------
set_lines(shape(S[3], 9), ['Italy'])
set_lines(shape(S[3], 10), ['Leonardo–Baykar and Rheinmetall', 'Italia in series production'])
set_lines(shape(S[3], 11), ['Find the design authority', 'behind each programme'])
set_lines(shape(S[3], 13), ['Spain  ·  Portugal'])
set_lines(shape(S[3], 14), ['Tarsis deliveries from 2028;', 'TEKEVER builds AR5 in the UK'])
note_replace(S[3], 'Portugal is second: TEKEVER is building the AR5 for the British Army at its UK sites, so first I would find out whether the decision sits in the UK or in Portugal, and who already owns the account. Italy and Spain I would qualify rather than lead with.',
             'Italy is second: Leonardo–Baykar and Rheinmetall Italia are in series production, so the first job is finding the design authority behind each programme. Spain and Portugal I would qualify rather than lead with, even though they are my home ground: TEKEVER builds the AR5 at its UK sites and uses its own communications, so the decision may not sit in Portugal at all.')
note_replace(S[3], 'What I propose: Proposed order: France, then Portugal, then qualify Italy and Spain.',
             'What I propose: Proposed order: France, then Italy, then qualify Spain and Portugal.')

# ---------- Slide 4: drop the redundant "Lead with" row, close the gap ----------
delete(S[4], 18)
delete(S[4], 19)
move_y(S[4], 20, 3.85); move_y(S[4], 21, 3.85)
move_y(S[4], 22, 4.25); move_y(S[4], 23, 4.25)

# ---------- Slide 5: one line of personal proof ----------
clone_below(S[5], 15, 4.78, ['How I opened Portugal’s grid operator from zero: a partner who held the relationship.'], bold=True)
note_replace(S[5], 'After a first win, I would ask about the next platform and whether the customer will act as a reference.',
             'After a first win, I would ask about the next platform and whether the customer will act as a reference. '
             'It is the route I used to open Portugal’s grid operator from zero: a partner already held the relationship and had a reason to hand it over; I built the numbers, secured the permit and we won the client’s tender.')

# ---------- Slide 6: industry language on the forecast ----------
set_lines(shape(S[6], 14), ['Radios × units', '× price = forecast'])
set_lines(shape(S[6], 27), ['A framework ceiling is not an order; an unweighted pipeline is not a forecast. Stage names would follow Doodle’s CRM.'])
note_replace(S[6], 'If a customer misses the agreed step, I move the line down or stop it, and I say so early.',
             'If a customer misses the agreed step, I move the line down or stop it, and I say so early. '
             'Listed defence suppliers report it the same way: Astronics has a $215 million Army framework and counts the $44.7 million first order, not the ceiling.')

# ---------- Slide 7: private-label channel ----------
clone_below(S[7], 16, 4.80, ['In Europe the winnable channel is often private-label: the radio inside a partner’s own product.'])
note_replace(S[7], 'A distributor that only holds stock can still help with fulfilment; I would just not count its stock as demand.',
             'A distributor that only holds stock can still help with fulfilment; I would just not count its stock as demand. '
             'In Europe the most valuable partner is often private-label: Solace sells MultiLink, Broadcast Solutions sells meshLINK — the deal is the radio inside their product. And every partner gets an export-control check before the first call: one candidate I found lists a Russia office.')

# ---------- Slides 8–10: Inherit / Do / Deliver, three bullets, proposed numbers ----------
for i in (8, 9, 10):
    s = S[i]
    set_lines(shape(s, 7), ['Inherit'])
    set_lines(shape(s, 15), ['Do'])
    set_lines(shape(s, 23), ['Deliver'])
    for sid in (8, 16, 24):
        delete(s, sid)

# Slide 8
s = S[8]
delete(s, 13); move_y(s, 14, 3.35)
set_lines(shape(s, 19), ['Existing: urgent risks, what’s next'])
set_lines(shape(s, 20), ['New: reconcile my map with CRM'])
set_lines(shape(s, 21), ['New: first programme calls'])
delete(s, 22)
set_lines(shape(s, 27), ['20 priority accounts (proposed)'])
set_lines(shape(s, 28), ['Each with an owner and a date'])
set_lines(shape(s, 29), ['Stage criteria agreed'])
delete(s, 30)
set_lines(shape(s, 33), ['Around 20 agreed priorities, each with an owner, a customer action and a date.'])
note_replace(s, 'At the day-thirty review I would bring agreed account priorities, each with an owner, a customer action and a date.',
             'At the day-thirty review I would bring around twenty agreed account priorities — a proposal, calibrated to your book — each with an owner, a customer action and a date.')

# Slide 9
s = S[9]
delete(s, 13); move_y(s, 14, 3.35)
set_lines(shape(s, 19), ['Existing: expansion, delivery'])
set_lines(shape(s, 20), ['New: sponsor, criteria, date'])
set_lines(shape(s, 21), ['New: partner on a named deal'])
delete(s, 22)
set_lines(shape(s, 27), ['4–6 live (proposed, sized to SE)'])
set_lines(shape(s, 28), ['An owner and a review date'])
set_lines(shape(s, 29), ['Blockers, with a recommendation'])
delete(s, 30)
set_lines(shape(s, 33), ['4–6 evaluations, each with a customer commitment behind it.'])
note_replace(s, 'At day sixty every active evaluation should have a customer commitment behind it.',
             'At day sixty I would expect four to six active evaluations — my proposal, to confirm against real Sales Engineering capacity — each with a customer commitment behind it.')
note_replace(s, 'No fixed number of evaluations and no assumption about one engineer’s capacity.',
             'The 4–6 is a proposal to confirm against actual Sales Engineering capacity, not an assumption about one engineer.')

# Slide 10
s = S[10]
delete(s, 12); move_y(s, 13, 3.06); move_y(s, 14, 3.35)
set_lines(shape(s, 19), ['Existing: expansion and orders'])
set_lines(shape(s, 20), ['New: next sourcing step'])
set_lines(shape(s, 21), ['New: stop, with the reason'])
delete(s, 22)
delete(s, 30)
note_replace(s, 'I would not assume that any design-in becomes revenue within ninety days.',
             'I would not assume that any design-in becomes revenue within ninety days. Even Astronics, selected on the Army’s MV-75, is in a multi-year development phase — about $100 million of work from 2024 to mid-2027 — before production content flows.')

# ---------- Slide 11: name Germany ----------
set_lines(shape(S[11], 10), ['BEYOND CORE'])
set_lines(shape(S[11], 11), ['Germany: named accounts or support, agreed with Andy.', 'Elsewhere, only where I add access a customer lacks.'])
note_replace(S[11], 'Outside the core I would take named accounts only where I add access or continuity a customer lacks today,',
             'Germany I name explicitly: it holds most of the evidenced volume and your Sales Engineer hire is there, so I would take named accounts or support there as Andy decides. Elsewhere I would take named accounts only where I add access or continuity a customer lacks today,')

p.save(OUT)
print('saved', OUT)
